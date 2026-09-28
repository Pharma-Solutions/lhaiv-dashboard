# Partner-confirmation matcher: OptiSource 165 partners -> licensed facilities
# Runs on Mark's machine (py -3.12); reads enriched Master Data; writes worklist to _working_docs.
import pandas as pd, re, glob, os, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

MD = r"C:\Users\MarkFulton\OneDrive - Pharma Solutions USA, Inc\All Company - Documents\LighthouseAI - Product\LighthouseAI Verified\Master Data"
VC = r"C:\Verified\Claude outputs\vendor_coverage.csv"
DEMO = {"KY","LA","MI","NC","OH","PA","TX","WA","PR","AL","OR"}  # member + sell-into

SUFFIX = re.compile(r"\b(inc|incorporated|llc|l l c|corp|corporation|co|company|ltd|limited|lp|llp|plc|pllc|pc|usa|us|na|the)\b", re.I)
GENERIC = re.compile(r"\b(pharmaceuticals?|pharma|laboratories|labs?|healthcare|health|sciences?|therapeutics|biotech|medical|products?|group|holdings?|international|industries|manufacturing|distribution|distributors?|solutions?)\b", re.I)
def norm(s):
    s=str(s).lower()
    s=re.sub(r"\(.*?\)"," ",s)
    s=re.sub(r"[^a-z0-9 ]"," ",s)
    s=SUFFIX.sub(" ",s)
    return re.sub(r"\s+"," ",s).strip()
def stem(s):  # distinctive core: drop generic industry words too
    s=GENERIC.sub(" ",norm(s))
    return re.sub(r"\s+"," ",s).strip()

def pick(cols, pats, avoid=()):
    for p in pats:
        for c in cols:
            cl=c.lower()
            if re.search(p, cl) and not any(re.search(a,cl) for a in avoid): return c
    return None
def namecols(cols):
    # ordered business/org name candidates (try each per row); person-name-only fields excluded
    pats=[r"^facility name$",r"servicedfacility_businessname",r"pharmacy_businessname",r"^businessname$",r"^business name$",
          r"doingbusinessas",r"^dba$",r"location_name",r"entityname",r"^sort name$",r"^business$",r"^name$",r"combined_name",
          r"full_name",r"^fullname$",r"practitionername",r"org"]
    avoid=[r"first",r"last",r"middle",r"owner",r"responsible",r"contact",r"file",r"user",r"representative",r"person"]
    out=[]
    for p in pats:
        for c in cols:
            cl=c.lower()
            if re.search(p,cl) and not any(re.search(a,cl) for a in avoid) and c not in out: out.append(c)
    return out
def detect(cols):
    return {
     "name": (namecols(cols)[0] if namecols(cols) else None),
     "names": namecols(cols),
     "lic": pick(cols,[r"license\s*#",r"license\s*number",r"licensenumber",r"license_no",r"^license no$",r"certificate_no",r"credentialnumber",r"credential number",r"cdspermit",r"lpr_num",r"^fnum$",r"licensepermitid$",r"lic_number",r"^license$"]),
     "type": pick(cols,[r"license sub type",r"license\s*type",r"licensetype",r"registration_type",r"facilitytype",r"factype",r"^type$",r"credentialtype",r"licensetypedescription",r"^class$",r"subcategory",r"profession"]),
     "status": pick(cols,[r"license status",r"licensestatus",r"applicant_status",r"^status$",r"credential.*status"]),
     "addr": pick(cols,[r"^address 1$",r"address1",r"^address$",r"mailing address 1",r"primary address 1",r"phys_address1",r"facilityaddress",r"streetaddress",r"addr 1",r"addr_line_1",r"ba_address",r"mailingstreetaddress",r"address line 1",r"^adress$",r"sort name / address"]),
     "issue": pick(cols,[r"issue date",r"dateissued",r"issue_date",r"effective_date",r"originallicensedate",r"licensefirstissuedate",r"current_issue_date",r"firstissuedate",r"license date",r"orig_license_date"]),
     "exp": pick(cols,[r"expire",r"expiration",r"dateexpired",r"expiration_date",r"licenseexpirationdate",r"current_expiration_date",r"license expired"]),
     "resnon": pick(cols,[r"resident_nonresident"]),
    }

# ---- partners ----
v=pd.read_csv(VC,dtype=str,keep_default_na=False)
partners=[(r["Account Name"].strip(), r.get("Priority","").strip()) for _,r in v.iterrows() if r["Account Name"].strip()]
pindex={}  # stem -> list of (partner, norm)
for p,_ in partners:
    pindex.setdefault(stem(p),[]).append((p,norm(p)))
print(f"partners: {len(partners)}")

# ---- scan facilities ----
matches=[]  # partner, priority, state, facility, lic, type, status, resnon, completeness, confidence, file
prio={p:pr for p,pr in partners}
files=[f for f in glob.glob(os.path.join(MD,"*_enriched.csv"))]
def st_of(fn):
    b=os.path.basename(fn)
    return b[:2].upper() if (len(b)>2 and (b[2] in " -_")) else b[:2].upper()
scanned=0
for f in files:
    st=st_of(f)
    if st not in DEMO: continue
    try:
        df=pd.read_csv(f,dtype=str,keep_default_na=False)
    except Exception as e:
        print("skip",os.path.basename(f),e); continue
    scanned+=1
    m=detect(list(df.columns))
    namecands=m["names"]
    if not namecands:
        print("NO NAME COL:",os.path.basename(f),list(df.columns)[:6]); continue
    fields7=[m["name"],m["addr"],m["type"],m["lic"],m["issue"],m["exp"],m["status"]]
    for _,row in df.iterrows():
        hit=None; conf=""; fname=""
        for nc in namecands:
            val=str(row.get(nc,"")).strip()
            if not val: continue
            fs=stem(val); fn=norm(val)
            cands=pindex.get(fs,[])
            for p,pn in cands:
                if pn==fn: hit=p; conf="exact"; fname=val; break
            if hit: break
            for p,pn in cands:
                if len(pn.split())>=2 and (pn in fn or fn in pn): hit=p; conf="strong (stem+substr)"; fname=val; break
            if hit: break
        if hit:
            present=sum(1 for c in fields7 if c and str(row.get(c,"")).strip())
            comp=round(present/7*100)
            matches.append({"Partner":hit,"Priority":prio.get(hit,""),"State":st,
                "Matched facility":fname,"License #":row.get(m["lic"],"") if m["lic"] else "",
                "License type":row.get(m["type"],"") if m["type"] else "",
                "Status":row.get(m["status"],"") if m["status"] else "",
                "Res/nonres":row.get(m["resnon"],"") if m["resnon"] else "",
                "Completeness %":comp,"Confidence":conf,"Source file":os.path.basename(f)})
print(f"files scanned: {scanned} | raw matches: {len(matches)}")
M=pd.DataFrame(matches)
# dedupe identical (partner,state,facility,lic)
if len(M): M=M.drop_duplicates(subset=["Partner","State","Matched facility","License #"])
print(f"deduped matches: {len(M)}")

# coverage per partner
cov=[]
for p,pr in partners:
    sub=M[M["Partner"]==p] if len(M) else M
    sts=sorted(sub["State"].unique()) if len(sub) else []
    cov.append({"Partner":p,"Priority":pr,"States matched":len(sts),"Facilities":len(sub),
                "States":", ".join(sts),"Any match?":"Yes" if len(sub) else "No"})
C=pd.DataFrame(cov)
matched_partners=(C["Any match?"]=="Yes").sum()
print(f"partners with >=1 facility: {matched_partners} of {len(partners)}")

# ---- workbook ----
NAVY="1F3A5F";HDR=Font(name="Arial",bold=True,color="FFFFFF",size=9);TITLE=Font(name="Arial",bold=True,size=13,color="1F3A5F")
SUB=Font(name="Arial",italic=True,size=9,color="555555");BODY=Font(name="Arial",size=9)
NF=PatternFill("solid",fgColor=NAVY);GREEN=PatternFill("solid",fgColor="C6EFCE");RED=PatternFill("solid",fgColor="FCE4E4");YEL=PatternFill("solid",fgColor="FFF2CC")
thin=Side(style="thin",color="D9D9D9");BD=Border(left=thin,right=thin,top=thin,bottom=thin);WRAP=Alignment(wrap_text=True,vertical="top")
wb=openpyxl.Workbook()
def sheet(ws,title,sub,cols,widths,rows,fillfn=None):
    ws["A1"]=title;ws["A1"].font=TITLE;ws["A2"]=sub;ws["A2"].font=SUB
    ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=len(cols));ws.row_dimensions[2].height=30
    for j,(h,w) in enumerate(zip(cols,widths),1):
        c=ws.cell(4,j,h);c.font=HDR;c.fill=NF;c.alignment=Alignment(wrap_text=True,vertical="center");c.border=BD;ws.column_dimensions[get_column_letter(j)].width=w
    for i,rw in enumerate(rows,5):
        for j,v in enumerate(rw,1):
            c=ws.cell(i,j,v);c.font=BODY;c.alignment=WRAP;c.border=BD
        if fillfn: fillfn(ws,i,rw)
    ws.freeze_panes="A5"

# Tab 1 coverage summary
ws1=wb.active;ws1.title="Partner coverage"
Cs=C.sort_values(["Any match?","States matched","Facilities"],ascending=[True,False,False])
def f1(ws,i,rw):
    ws.cell(i,6).fill = GREEN if rw[5]=="Yes" else RED
sheet(ws1,"OptiSource 165 — partner facility coverage",
      f"{matched_partners} of {len(partners)} partners matched to >=1 licensed facility across {len(DEMO)} demo states (KY LA MI NC OH PA TX WA PR AL OR). Normalized-name match; propose/review (human-gated). Red = no facility match found.",
      ["Partner","Priority","States matched","Facilities","States","Any match?"],[34,10,15,12,34,12],
      [[r["Partner"],r["Priority"],r["States matched"],r["Facilities"],r["States"],r["Any match?"]] for _,r in Cs.iterrows()],f1)
# Tab 2 facility matches
ws2=wb.create_sheet("Facility matches")
Ms=M.sort_values(["Partner","State"]) if len(M) else M
def f2(ws,i,rw):
    try: comp=int(rw[8])
    except: comp=0
    ws.cell(i,9).fill = GREEN if comp>=85 else (YEL if comp>=60 else RED)
sheet(ws2,"Facility-level matches (per partner -> facility)",
      "Each matched licensed facility with license detail + completeness %. Completeness = fields present of 7 (name, address, type, license #, issue date, expiration date, status). Confidence: exact vs strong(stem+substr). Verify before writing confirmations (entity merges human-gated).",
      ["Partner","Priority","State","Matched facility","License #","License type","Status","Res/nonres","Completeness %","Confidence","Source file"],
      [26,9,6,34,16,26,14,12,13,20,34],
      [[r["Partner"],r["Priority"],r["State"],r["Matched facility"],r["License #"],r["License type"],r["Status"],r["Res/nonres"],r["Completeness %"],r["Confidence"],r["Source file"]] for _,r in Ms.iterrows()],f2)
# Tab 3 unmatched
un=C[C["Any match?"]=="No"].sort_values("Priority")
ws3=wb.create_sheet("Unmatched partners")
sheet(ws3,f"Unmatched partners ({len(un)}) — no facility found in demo states",
      "No normalized-name facility match in the 11 demo states. Could be: not licensed in a demo state, a name variant/DBA the match missed, or a data gap. Review priority-High first.",
      ["Partner","Priority"],[40,12],[[r["Partner"],r["Priority"]] for _,r in un.iterrows()])
# Tab 4 method
ws4=wb.create_sheet("Method")
ws4["A1"]="Method & caveats";ws4["A1"].font=TITLE;ws4.column_dimensions["A"].width=118
notes=[
 "Partner list: OptiSource 165 (vendor_coverage.csv). NOTE - reconcile against the authoritative 'OptiSource Manufacturers.xlsx' (customer-provided) and the deployment entities '6 partners' sheet before finalizing.",
 "Facilities: all *_enriched.csv in Master Data for the 11 demo states (KY, LA, MI, NC, OH, PA, TX, WA, PR, AL, OR). KY spans 8 sub-lists (mfr/wholesaler/TPL/outsourcer/pharmacy/etc.).",
 "Column roles (name/license #/type/status/address/res-nonres) auto-detected per file by header pattern (schemas are heterogeneous).",
 "Match rule: normalize names (drop corporate suffixes + generic industry words to a distinctive stem); exact normalized match, or strong stem+substring with >=2 partner tokens. Conservative to limit false positives; a missed match is safer than a wrong confirmation.",
 "Completeness % = of the 7 MVP fields present per facility. Drives the demo's license-profile completeness metric.",
 "PROPOSE/REVIEW only — entity merges and partner<->facility confirmations are human-in-the-loop gates. This worklist is the input to the Results surface, not an auto-write.",
 "Unmatched partners are not necessarily gaps — many suppliers won't be licensed in all 11 states; the coverage tab shows breadth per partner.",
]
r=3
for n in notes:
    c=ws4.cell(r,1,n);c.font=BODY;c.alignment=WRAP;ws4.row_dimensions[r].height=(30 if len(n)<115 else 46);r+=1
out=r"C:\Verified\_working_docs\Verified_Partner_Confirmation_Worklist_20260924.xlsx"
wb.save(out); print("saved",out)
