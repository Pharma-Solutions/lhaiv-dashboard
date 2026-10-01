#!/usr/bin/env python3
"""Generic intake engine: emailed board workbook -> canonical roster CSV.

One engine, many jurisdictions. Everything state-specific is data in
email_sources.py; this file holds only the transformation and the guards, so a
defect fixed here is fixed for every state at once and a new state is a registry
entry rather than another bespoke script to audit.

GUARDS (each exists because something went wrong without it)
  * header row is DETECTED, never assumed - MD's headers start on row 3, and a
    naive read returns a frame with blank column names and two junk records
    without raising.
  * an unmapped worksheet is FATAL - the board partitions by sheet, so a sheet we
    do not recognise is licences we would silently drop.
  * a sheet below half its expected count is FATAL; any other deviation is loud.
    Never overwrite a good file with a short one.
  * whole-row de-dupe only. Collapsing on a natural key would merge genuinely
    distinct records that share a licence number (the Colorado lesson).
  * zip padding is COUNTED only when it actually changes the value, so the
    reported figure is reconstructions performed, not integers seen.
  * a status outside the registered enum is kept verbatim and flagged, never
    mapped to the nearest known value.

Usage:  python email_intake.py WY GA MD        (or --all)
"""
import argparse
import collections
import csv
import datetime as dt
import hashlib
import json
import os
import re
import sys

import openpyxl

import email_sources as ES

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

DATA = r"C:\Verified\data"


def incoming_dir(state):
    return os.path.join(DATA, state, "_incoming")


def squash(s):
    return re.sub(r"\s+", " ", str(s or "")).strip()


def txt(v):
    return "" if v is None else str(v).strip()


def sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


class Intake(object):
    def __init__(self, spec, log):
        self.spec = spec
        self.log = log
        self.stats = collections.Counter()
        self.files = []

    # ---------------------------------------------------------------- helpers
    def norm_date(self, v, where):
        if v is None or v == "":
            return ""
        if isinstance(v, (dt.datetime, dt.date)):
            return v.strftime("%Y-%m-%d")
        self.stats["date_not_datetime"] += 1
        if self.stats["date_not_datetime"] <= 5:
            self.log("   !! non-datetime date in %s: %r (kept verbatim)" % (where, v))
        return str(v).strip()

    def norm_zip(self, v):
        if v is None:
            return ""
        s = str(v).strip()
        if not s:
            return ""
        if isinstance(v, int) or re.fullmatch(r"\d{1,4}", s):
            padded = s.zfill(5)
            if padded != s:                 # count reconstructions, not integers
                self.stats["zip_zero_padded"] += 1
            return padded
        return s

    def header_row(self, ws, scan=12):
        for r in range(1, min(scan, ws.max_row or 1) + 1):
            vals = [c.value for c in ws[r]]
            if sum(1 for v in vals if isinstance(v, str) and v.strip()) >= 3:
                return r, [squash(v).lower() for v in vals]
        raise SystemExit("[%s] FATAL: no header row in sheet %r" % (self.spec.state, ws.title))

    def resolve(self, hdr):
        """canonical field -> column index, using this source's aliases."""
        ix = {}
        for canon, names in self.spec.aliases.items():
            found = None
            for n in names:
                if n in hdr:
                    found = hdr.index(n)
                    break
            ix[canon] = found
        return ix

    # ------------------------------------------------------------------ sheet
    def read_sheet(self, ws, lic_type, label, src_file):
        hr, hdr = self.header_row(ws)
        ix = self.resolve(hdr)
        required = ["license_number", "facility_name"]
        missing = [c for c in required if ix.get(c) is None]
        if missing:
            raise SystemExit("[%s] FATAL: %s missing required column(s) %s; header=%s"
                             % (self.spec.state, label, missing, [h for h in hdr if h]))
        if lic_type is ES.TYPE_FROM_COLUMN and ix.get("license_type") is None:
            raise SystemExit("[%s] FATAL: %s declares TYPE_FROM_COLUMN but no type column "
                             "matched; header=%s" % (self.spec.state, label, [h for h in hdr if h]))

        rows = []
        for r in ws.iter_rows(min_row=hr + 1, values_only=True):
            if not r:
                continue

            def g(c):
                i = ix.get(c)
                return r[i] if (i is not None and i < len(r)) else None

            lic = txt(g("license_number"))
            if not lic:
                continue                                  # trailing blank rows
            status = txt(g("license_status"))
            if self.spec.known_status is not None and status and status not in self.spec.known_status:
                self.stats["status_drift"] += 1
                self.log("   !! unmapped status %r in %s (kept verbatim)" % (status, label))
            rows.append({
                "facility_name":    txt(g("facility_name")),
                "license_number":   lic,
                "license_type":     txt(g("license_type")) if lic_type is ES.TYPE_FROM_COLUMN else lic_type,
                "license_status":   status,
                "business_activity": txt(g("business_activity")),
                "issue_date":       self.norm_date(g("issue_date"), label),
                "expiration_date":  self.norm_date(g("expiration_date"), label),
                "address_line1":    txt(g("address_line1")),
                "address_line2":    txt(g("address_line2")),
                "address_city":     txt(g("address_city")),
                "address_county":   txt(g("address_county")),
                "address_state":    txt(g("address_state")).upper(),
                "address_zip":      self.norm_zip(g("address_zip")),
                "jurisdiction":     self.spec.state,
                "__source_sheet":   ws.title,
                "__source_file":    src_file,
                "__source_email":   self.spec.source_string(),
                "__retrieved":      self.spec.email["received"],
            })
        return rows

    # ------------------------------------------------------------------- main
    def run(self, out_dir, datestamp):
        sp = self.spec
        frames = []
        self.log("=" * 76)
        self.log("%s  %s" % (sp.state, sp.agency))
        self.log("  %s" % sp.source_string())
        self.log("=" * 76)

        for wbspec in sp.workbooks:
            path = os.path.join(incoming_dir(sp.state), wbspec.filename)
            if not os.path.exists(path):
                raise SystemExit("[%s] FATAL: missing source workbook %s" % (sp.state, path))
            digest = sha256(path)
            self.files.append({"file": wbspec.filename, "sha256": digest,
                               "bytes": os.path.getsize(path)})
            self.log("  %s" % wbspec.filename)
            self.log("    sha256 %s" % digest)

            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            if wbspec.sheets is ES.TYPE_FROM_COLUMN:
                sheets = {s: ES.TYPE_FROM_COLUMN for s in wb.sheetnames}
            else:
                unmapped = [s for s in wb.sheetnames if s not in wbspec.sheets]
                if unmapped:
                    raise SystemExit("[%s] FATAL: unmapped sheet(s) %s in %s - refusing to "
                                     "drop them silently" % (sp.state, unmapped, wbspec.filename))
                sheets = wbspec.sheets

            for sheet in wb.sheetnames:
                label = "%s/%s" % (wbspec.filename, sheet)
                rows = self.read_sheet(wb[sheet], sheets[sheet], label, wbspec.filename)
                exp = wbspec.expected_rows.get(sheet)
                tag = ""
                if exp is not None:
                    tag = "  (expected %d%s)" % (exp, "" if len(rows) == exp else "  <-- DELTA")
                    if len(rows) != exp:
                        self.stats["count_delta"] += 1
                    if len(rows) < exp * 0.5:
                        raise SystemExit("[%s] FATAL: %s gave %d rows vs %d expected - refusing "
                                         "to write a short file" % (sp.state, label, len(rows), exp))
                self.log("    %-16s %6d rows%s" % (sheet, len(rows), tag))
                frames.extend(rows)
            wb.close()

        if not frames:
            raise SystemExit("[%s] FATAL: no rows parsed - refusing to write an empty file" % sp.state)

        # lossless whole-row de-dupe
        seen, deduped = set(), []
        for r in frames:
            k = tuple(r[c] for c in ES.CANONICAL)
            if k in seen:
                self.stats["whole_row_dup"] += 1
                self.log("    whole-row duplicate dropped: %s %r"
                         % (r["license_number"], r["facility_name"]))
                continue
            seen.add(k)
            deduped.append(r)
        frames = deduped

        # near-duplicate tripwire: documented ones are kept, new ones are surfaced
        within = collections.Counter((r["license_number"], r["license_type"]) for r in frames)
        repeats = {lic for (lic, _t), n in within.items() if n > 1}
        for lic in sorted(repeats - sp.documented_near_dups):
            self.stats["new_near_dup"] += 1
            self.log("    !! NEW near-duplicate not in the documented set: %s" % lic)
        self.stats["near_dup_known"] = len(repeats & sp.documented_near_dups)

        out = os.path.join(out_dir, sp.out_name(datestamp))
        with open(out, "w", encoding="utf-8-sig", newline="") as f:
            w = csv.DictWriter(f, fieldnames=ES.CANONICAL)
            w.writeheader()
            w.writerows(frames)

        prov = {
            "state": sp.state, "agency": sp.agency, "scope": sp.scope,
            "email": sp.email, "source_files": self.files,
            "output": os.path.basename(out), "output_sha256": sha256(out),
            "rows": len(frames),
            "distinct_license_numbers": len({r["license_number"] for r in frames}),
            "by_type": dict(collections.Counter(r["license_type"] for r in frames)),
            "by_status": dict(collections.Counter(r["license_status"] for r in frames)),
            "stats": dict(self.stats),
            "notes": sp.notes,
            "built": dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        }
        pp = out[:-4] + ".provenance.json"
        json.dump(prov, open(pp, "w", encoding="utf-8"), indent=1)

        self.log("")
        self.log("  wrote %s" % out)
        self.log("  rows %s   distinct licence numbers %s"
                 % (format(len(frames), ","), format(prov["distinct_license_numbers"], ",")))
        for t, n in collections.Counter(r["license_type"] for r in frames).most_common():
            self.log("     %-42s %6d" % (t, n))
        self.log("  status        : %s" % (prov["by_status"] if sp.known_status is not None
                                           else "(source carries no status column)"))
        self.log("  zip padded    : %d" % self.stats["zip_zero_padded"])
        self.log("  date anomalies: %d" % self.stats["date_not_datetime"])
        self.log("  status drift  : %d" % self.stats["status_drift"])
        self.log("  count deltas  : %d" % self.stats["count_delta"])
        self.log("  whole-row dup : %d (lossless)" % self.stats["whole_row_dup"])
        self.log("  near-dup known: %d   NEW: %d"
                 % (self.stats["near_dup_known"], self.stats["new_near_dup"]))
        self.log("  provenance -> %s" % os.path.basename(pp))
        self.log("")
        return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("states", nargs="*", help="e.g. WY GA MD")
    ap.add_argument("--all", action="store_true")
    ap.add_argument("--date", default=dt.date.today().strftime("%Y%m%d"))
    a = ap.parse_args()
    keys = sorted(ES.ALL) if a.all or not a.states else [s.upper() for s in a.states]
    unknown = [k for k in keys if k not in ES.ALL]
    if unknown:
        raise SystemExit("unknown source(s) %s; registered: %s" % (unknown, sorted(ES.ALL)))

    def log(m=""):
        print(m, flush=True)

    for k in keys:
        Intake(ES.ALL[k], log).run(incoming_dir(k), a.date)


if __name__ == "__main__":
    main()
