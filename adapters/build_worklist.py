#!/usr/bin/env python3
"""Rebuild LHAIV_Source_PressureTest_Worklist.xlsx from the master tracker.

The harness zip shipped without the worklist, so we regenerate it in the exact
schema probe_sources.py expects: sheet "Source Pressure-Test", header on row 6
(header=5), columns whose captions contain State / Agency / URL / Bucket, plus
the yellow verdict columns a human fills in.
"""
import warnings; warnings.filterwarnings("ignore")
import re
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

SRC = "lhaiv-dashboard-pipeline/tracker.xlsx"
OUT = "LHAIV_Source_PressureTest_Worklist.xlsx"

d = pd.read_excel(SRC, sheet_name="Tracker simplified", header=1, dtype=str)
d.columns = [str(c).strip() for c in d.columns]
d = d[d["State"].notna()].fillna("")
ins = d[d["Outcome"].str.strip() != "Out of Scope"].copy()   # -> 76 in-scope

HOW  = "How to Download + What's in the Download"
PROV = "How is the data provided?"

def bucket(r):
    prov, out = r[PROV].strip(), r["Outcome"].strip()
    if prov == "File Download" and out == "Obtained - Free":
        return "A - bulk file already obtained (free)"
    if out == "Obtained - Purchased":
        return "C - purchased list already obtained"
    if prov == "File Download":
        return "D - bulk file channel identified, not yet pulled"
    if prov == "Published roster PDF":
        return "D - published roster identified"
    if "request" in prov.lower():
        return "B - records/roster request channel identified"
    if prov == "Query Feature":
        return "E - portal/query"
    return "F - undetermined"

ins["Bucket"] = ins.apply(bucket, axis=1)


# --- URL hygiene -----------------------------------------------------------
# The tracker stores prose inside some URL cells ("https://x (see also: y)"),
# which the probe then fetches verbatim and fails on. Split at the first space
# and keep the prose as a note. A few rows also point at a landing page when the
# data-bearing page is the one named in the prose - override those explicitly.
URL_OVERRIDE = {
    # Guam: the landing page is a stub; the paginated 1,032-licensee directory is here.
    ("Guam", "Guam Board of Examiners for Pharmacy"): "https://web.guamhplo.org/directory/gbep",
}

def clean_url(state, agency, raw):
    raw = (raw or "").strip()
    ov = URL_OVERRIDE.get((state.strip(), agency.strip()))
    if ov:
        return ov, f"tracker cell also read: {raw}"
    if not raw.lower().startswith("http"):
        return "", raw          # prose-only cell -> no URL, keep prose as note
    url = re.split(r"[\s(]", raw, 1)[0].rstrip(".,;")
    extra = raw[len(url):].strip(" ()") if len(raw) > len(url) else ""
    return url, extra

rows = []
for _, r in ins.iterrows():
    prior = " | ".join(x for x in [
        f"prov={r[PROV]}" if r[PROV] else "",
        f"outcome={r['Outcome']}" if r["Outcome"] else "",
        f"fee={r['Fee Amount ($)']}" if str(r["Fee Amount ($)"]).strip() not in ("", "0") else "",
        f"prior-Hermai={r['Hermai Recommended? (Y/N)']}" if r["Hermai Recommended? (Y/N)"] else "",
    ] if x)
    url, url_note = clean_url(r["State"], r["Agency"], r["Agency URL"])
    if url_note:
        prior = (prior + " | " if prior else "") + f"URL-note: {url_note}"
    rows.append({
        "State": r["State"].strip(),
        "Agency": r["Agency"].strip(),
        "Source URL": url,
        "Bucket": r["Bucket"],
        "Prior finding (from tracker)": prior,
        "Prior how-to / notes": (r[HOW].strip() + (" // " + r["Mark-Notes"].strip() if r["Mark-Notes"].strip() else "")).strip(),
        "Downloadable?": "", "Tier": "", "Access mechanism": "",
        "Platform": "", "Confirmed-by": "", "Date": "", "Notes": "",
    })

df = pd.DataFrame(rows).sort_values(["Bucket", "State", "Agency"], kind="stable")

wb = Workbook(); ws = wb.active; ws.title = "Source Pressure-Test"
ws["A1"] = "LHAIV Verified - Source Pressure-Test Worklist"
ws["A1"].font = Font(bold=True, size=14)
ws["A2"] = ("Goal: for every in-scope agency, prove whether full-population data is downloadable/extractable "
            "WITHOUT the paid scraper (Hermai), and by what mechanism.")
ws["A3"] = ("Tiers - 1: bulk download/open-data/API (automate) | 2: purchase or records/roster request (human gate) | "
            "3: enumerable no-CAPTCHA portal (scrape, one adapter per platform) | 4: CAPTCHA or single-record only (Hermai candidate)")
ws["A4"] = ("Buckets - A/C: data already in hand | B/D: channel identified | E: portal/query (probe) | F: undetermined (probe). "
            "Rebuilt from tracker.xlsx 'Tracker simplified' (117 rows less 41 Out of Scope = 76 in scope).")
# row 5 intentionally blank; row 6 = header (probe_sources.py reads header=5)

cols = list(df.columns)
YELLOW = {"Downloadable?", "Tier", "Access mechanism", "Platform", "Confirmed-by", "Date", "Notes"}
fill_y = PatternFill("solid", fgColor="FFF2CC")
fill_h = PatternFill("solid", fgColor="D9E1F2")
for j, c in enumerate(cols, 1):
    cell = ws.cell(row=6, column=j, value=c)
    cell.font = Font(bold=True)
    cell.fill = fill_y if c in YELLOW else fill_h
    cell.alignment = Alignment(wrap_text=True, vertical="bottom")

for i, (_, r) in enumerate(df.iterrows(), start=7):
    for j, c in enumerate(cols, 1):
        cell = ws.cell(row=i, column=j, value=r[c])
        if c in YELLOW:
            cell.fill = fill_y

widths = {"State": 16, "Agency": 46, "Source URL": 62, "Bucket": 40,
          "Prior finding (from tracker)": 46, "Prior how-to / notes": 60,
          "Downloadable?": 22, "Tier": 8, "Access mechanism": 34,
          "Platform": 22, "Confirmed-by": 16, "Date": 12, "Notes": 60}
for j, c in enumerate(cols, 1):
    ws.column_dimensions[get_column_letter(j)].width = widths.get(c, 18)
ws.freeze_panes = "A7"
wb.save(OUT)

print(f"wrote {OUT}: {len(df)} in-scope sources")
print(df["Bucket"].value_counts().sort_index().to_string())
has = df["Source URL"].str.lower().str.startswith("http")
print(f"\nwith http URL: {has.sum()}   non-URL (manual review): {(~has).sum()}")
print("E+F count:", df["Bucket"].str.startswith(("E", "F")).sum())
