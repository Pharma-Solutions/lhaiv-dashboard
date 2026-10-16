#!/usr/bin/env python3
"""HI NED CDS establishment lists: three text PDFs -> one xlsx, one sheet per list.

WHY A SEPARATE STEP
  email_intake.py reads workbooks. Hawaii NED delivered PDFs, so this converts
  them, documented and re-runnable, and email_intake then treats the xlsx like any
  other board workbook. The first HI build (2026-10-14) used `pdftotext -layout`
  and split columns by character position. Layout mode collapses the gaps between
  columns, so it truncated street numbers ("4401 E. U.S. HIGHWAY 30" became
  "01 E. U.S. HIGHWAY"), let header text bleed into data, and could not separate
  street from city. Here every word carries its PDF x-coordinate, and is assigned
  to a column by the header's own x-position ON THAT PAGE. The column positions
  differ between the three files (Pharmacy's State/Zip sit ~17pt left of
  Distributor's), so a single set of bounds would be wrong for at least one.

LAYOUT FACTS THIS RELIES ON (measured 2026-10-16, each one asserted at runtime)
  * every page repeats the same 12-column header; labels wrap onto up to three
    lines and are BOTTOM-aligned on the line that holds "NED Number".
  * data cells are BOTTOM-aligned too: a wrapped name or street puts its earlier
    lines ABOVE the line carrying the NED number. A record is therefore every word
    below the previous NED number and down to and including its own.
  * nothing follows the last NED number on a page (no record splits across a page
    break, no footer). Any word found there is FATAL, not silently dropped.

Transcription only: values are kept verbatim except expiration dates. They are
written as real dates, and only after every value is proved to be DD-MM-YYYY.

Usage:  python hi_ned_pdf_to_xlsx.py SRC_DIR OUT_XLSX [--report JSON]
"""
import argparse
import collections
import datetime as dt
import hashlib
import json
import os
import re
import subprocess
import sys

import openpyxl
import pdfplumber

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

HEADER = ["NED Number", "Expiration Date", "Status", "Full Name", "Last Name",
          "First Name", "Middle Name", "Organization Name", "Business Address Street",
          "Business Address City", "Business Address State", "Business Address Zip"]

# file -> sheet name (= license_type in the SourceSpec). There is no type column; the
# type is which list the registrant is on.
LISTS = [("Manufacturer List.pdf", "Manufacturer"),
         ("Distributor List.pdf", "Distributor"),
         ("Pharmacy List.pdf", "Pharmacy")]

NED = re.compile(r"E\d{5}")
PHRASE_GAP = 5.0      # words inside one header label are ~2pt apart, labels >= 7pt
TOL = 1.0


def sha256(p):
    h = hashlib.sha256()
    with open(p, "rb") as f:
        for b in iter(lambda: f.read(1 << 20), b""):
            h.update(b)
    return h.hexdigest()


def header_columns(words, where):
    """-> (column start x0 list, header-bottom top). Asserts the 12 labels exactly."""
    neds = [w for w in words if w["text"] == "NED"]
    if len(neds) != 1:
        raise SystemExit("FATAL %s: expected one 'NED' header word, found %d" % (where, len(neds)))
    base = neds[0]["top"]
    band = [w for w in words if w["top"] <= base + TOL]
    bottom = sorted((w for w in band if abs(w["top"] - base) <= TOL), key=lambda w: w["x0"])
    phrases = []
    for w in bottom:
        if phrases and w["x0"] - phrases[-1]["x1"] < PHRASE_GAP:
            phrases[-1]["words"].append(w["text"])
            phrases[-1]["x1"] = w["x1"]
        else:
            phrases.append({"x0": w["x0"], "x1": w["x1"], "words": [w["text"]], "above": []})
    upper = [w for w in band if w["top"] < base - TOL]
    for w in upper:
        hit = [p for p in phrases if abs(p["x0"] - w["x0"]) <= TOL]
        if len(hit) != 1:
            raise SystemExit("FATAL %s: header word %r at x0=%.1f belongs to no label"
                             % (where, w["text"], w["x0"]))
        hit[0]["above"].append(w)
    labels = []
    for p in phrases:
        above = [w["text"] for w in sorted(p["above"], key=lambda w: w["top"])]
        labels.append(" ".join(above + p["words"]))
    if labels != HEADER:
        raise SystemExit("FATAL %s: header labels differ from the expected 12:\n  %s"
                         % (where, labels))
    return [p["x0"] for p in phrases], base


def column_of(x0, starts):
    i = None
    for j, s in enumerate(starts):
        if x0 >= s - TOL:
            i = j
    return i


def char_key(c):
    return (c["text"], round(c["x0"], 2), round(c["top"], 2))


def join_words(ws, pos, stream, stats, notes, where, ned, col):
    """Join a cell's words using the PDF's own character stream.

    The report generator wraps a cell two ways. At a word boundary it emits the space
    character and continues on the next line ("MCKESSON |CORPORATION"). When one token
    is wider than the cell it breaks mid-token with no space ("HEALTHWAREHOUSE.C|OM",
    "PHARMACEUTICALS,INC|."). So whether two words are joined with a space is read from
    the stream, not guessed from geometry: the char after a word's last char is either
    a space (join with a space) or the next word's first char (join with nothing).
    """
    out = ""
    for j, w in enumerate(ws):
        if j:
            prev = ws[j - 1]
            k = pos.get(char_key(prev["chars"][-1]))
            nxt = stream[k + 1] if k is not None and k + 1 < len(stream) else None
            if nxt is not None and nxt["text"] == " ":
                out += " "
                stats["join_space"] += 1
            elif nxt is not None and char_key(nxt) == char_key(w["chars"][0]):
                stats["join_nospace"] += 1
                if len(notes) < 60:
                    notes.append("no-space join %s %s %s: %r+%r"
                                 % (where, ned, HEADER[col], prev["text"], w["text"]))
            else:
                # The stream jumps elsewhere: no evidence either way. Space, and count it.
                out += " "
                stats["join_unproven"] += 1
                notes.append("UNPROVEN join %s %s %s: %r+%r"
                             % (where, ned, HEADER[col], prev["text"], w["text"]))
        out += w["text"]
    return out


def parse_pdf(path, stats, notes):
    recs = []
    with pdfplumber.open(path) as pdf:
        for pno, page in enumerate(pdf.pages, start=1):
            where = "%s p%d" % (os.path.basename(path), pno)
            words = page.extract_words(x_tolerance=1.5, keep_blank_chars=False, return_chars=True)
            stream = page.chars
            pos = {char_key(c): k for k, c in enumerate(stream)}
            starts, base = header_columns(words, where)
            stats["pages"] += 1
            data = [w for w in words if w["top"] > base + TOL]
            anchors = sorted((w for w in data if column_of(w["x0"], starts) == 0),
                             key=lambda w: w["top"])
            for a in anchors:
                if not NED.fullmatch(a["text"]):
                    raise SystemExit("FATAL %s: non-NED text %r in the NED Number column"
                                     % (where, a["text"]))
            if not anchors:
                raise SystemExit("FATAL %s: no NED numbers on the page" % where)
            cells = [collections.defaultdict(list) for _ in anchors]
            for w in data:
                col = column_of(w["x0"], starts)
                if col is None:
                    raise SystemExit("FATAL %s: word %r left of the first column" % (where, w["text"]))
                # A word whose right edge runs into the next column would mean its
                # cell overflowed or the bounds are wrong. Count, and show the first few.
                if col + 1 < len(starts) and w["x1"] > starts[col + 1] - TOL:
                    stats["straddle"] += 1
                    if stats["straddle"] <= 10:
                        notes.append("straddle %s: %r col=%s" % (where, w["text"], HEADER[col]))
                idx = next((i for i, a in enumerate(anchors) if a["top"] >= w["top"] - TOL), None)
                if idx is None:
                    raise SystemExit("FATAL %s: word %r below the last NED number - the "
                                     "bottom-alignment assumption is broken" % (where, w["text"]))
                cells[idx][col].append(w)
            for i, a in enumerate(anchors):
                row = []
                for col in range(len(HEADER)):
                    ws = sorted(cells[i][col], key=lambda w: (round(w["top"], 1), w["x0"]))
                    if len({round(w["top"], 1) for w in ws}) > 1:
                        stats["wrapped_cells"] += 1
                    row.append(join_words(ws, pos, stream, stats, notes, where, a["text"], col))
                # Fixed columns must sit on the anchor line alone; a wrap there would mean
                # a neighbouring record's text was captured.
                for col in (0, 1, 2, 10, 11):
                    if any(abs(w["top"] - a["top"]) > TOL for w in cells[i][col]):
                        raise SystemExit("FATAL %s: %s for %s spans lines"
                                         % (where, HEADER[col], a["text"]))
                recs.append(row)
    return recs


def raw_ned_count(path):
    """Independent axis: pdftotext (no -layout), a different extractor entirely."""
    out = subprocess.run(["pdftotext", "-enc", "UTF-8", path, "-"], capture_output=True,
                         check=True).stdout.decode("utf-8", "replace")
    return len(re.findall(r"\bE\d{5}\b", out)), out


def to_date(s):
    m = re.fullmatch(r"(\d{2})-(\d{2})-(\d{4})", s)
    if not m:
        raise SystemExit("FATAL: expiration %r is not DD-MM-YYYY" % s)
    d, mo, y = (int(g) for g in m.groups())
    return dt.date(y, mo, d)            # raises on a non-date


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("src_dir")
    ap.add_argument("out_xlsx")
    ap.add_argument("--report", default=None)
    a = ap.parse_args()

    stats, notes, report = collections.Counter(), [], {"files": [], "controls": {}}
    sheets = {}
    for fn, sheet in LISTS:
        p = os.path.join(a.src_dir, fn)
        if not os.path.exists(p):
            raise SystemExit("FATAL: missing %s" % p)
        before = dict(stats)
        recs = parse_pdf(p, stats, notes)
        raw_n, raw_text = raw_ned_count(p)
        hdr_hits = raw_text.count("NED Number")
        report["files"].append({
            "file": fn, "sheet": sheet, "sha256": sha256(p), "bytes": os.path.getsize(p),
            "pages": stats["pages"] - before.get("pages", 0),
            "rows_out": len(recs), "distinct_ned": len({r[0] for r in recs}),
            "raw_ned_tokens_pdftotext": raw_n,
            "raw_header_repeats": hdr_hits,
        })
        if raw_n != len(recs):
            raise SystemExit("FATAL %s: pdftotext finds %d NED tokens, parser emitted %d rows"
                             % (fn, raw_n, len(recs)))
        sheets[sheet] = recs

    allrecs = [r for s in sheets.values() for r in s]
    # Dates: prove the format before converting. A DD-MM-YYYY claim needs at least one
    # value whose first field exceeds 12, otherwise it is indistinguishable from MM-DD.
    exp = [r[1] for r in allrecs]
    first_gt12 = sum(1 for e in exp if re.fullmatch(r"\d{2}-\d{2}-\d{4}", e) and int(e[:2]) > 12)
    second_gt12 = sum(1 for e in exp if re.fullmatch(r"\d{2}-\d{2}-\d{4}", e) and int(e[3:5]) > 12)
    if not first_gt12 or second_gt12:
        raise SystemExit("FATAL: date order not proved DD-MM (first>12: %d, second>12: %d)"
                         % (first_gt12, second_gt12))
    report["controls"]["date_order"] = {"values": len(exp), "first_field_gt12": first_gt12,
                                        "second_field_gt12": second_gt12}

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sheet, recs in sheets.items():
        ws = wb.create_sheet(sheet)
        ws.append(HEADER)
        for r in recs:
            out = list(r)
            out[1] = to_date(r[1]) if r[1] else None
            ws.append(out)
            # Text, so a ZIP such as 02134 or 1234 is not coerced to a number.
            for c in (1, 12):
                ws.cell(row=ws.max_row, column=c).number_format = "@"
            ws.cell(row=ws.max_row, column=2).number_format = "yyyy-mm-dd"
    wb.save(a.out_xlsx)

    report["stats"] = dict(stats)
    report["notes"] = notes
    report["output"] = {"file": os.path.basename(a.out_xlsx), "sha256": sha256(a.out_xlsx)}
    report["built"] = dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    for f in report["files"]:
        print("  %-24s pages=%-3d rows=%-4d distinct=%-4d pdftotext NED tokens=%d"
              % (f["file"], f["pages"], f["rows_out"], f["distinct_ned"],
                 f["raw_ned_tokens_pdftotext"]))
    print("  wrapped cells=%d  straddling words=%d" % (stats["wrapped_cells"], stats["straddle"]))
    print("  word joins: space=%d  no-space (mid-token wrap)=%d  UNPROVEN=%d"
          % (stats["join_space"], stats["join_nospace"], stats["join_unproven"]))
    for n in notes:
        print("   ", n)
    print("  wrote %s" % a.out_xlsx)
    if a.report:
        json.dump(report, open(a.report, "w", encoding="utf-8"), indent=1)


if __name__ == "__main__":
    main()
