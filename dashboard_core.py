"""
Core transform + render for the LHAIV license-data dashboard.

extract(xlsx_path) -> dict     auto-detects the tracker schema and parses it
render(data, build_str) -> str injects data into the matching HTML template

Two schemas are supported:
  * "discovery"  — the original API Discovery Tracker (coverage-ladder model)
  * "simplified" — the newer License Data Tracker (Y/N availability + cost + Hermai)

The renderer picks the template by schema, so pointing the pipeline at a new
source file is a config change, not a code change. Templates (_template.txt,
_template_new.txt) are the single source of visual truth.
"""
import os
import re
import json
import numpy as np
import pandas as pd
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))

# ---------- schema detection ----------

DISCOVERY_SHEET = "API Discovery Tracker"


def _detect_schema(xlsx_path):
    """discovery = original API Discovery Tracker; simplified = License Data Tracker
    (any sheet with State + Agency + an Outcome/How-provided column — tolerant of header renames)."""
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        if DISCOVERY_SHEET in wb.sheetnames:
            return "discovery", None
        for name in wb.sheetnames:
            ws = wb[name]
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                low = {str(c).strip().lower() for c in row if c is not None}
                if "state" in low and "agency" in low and (
                        "outcome" in low or any("provided" in v for v in low)):
                    return "simplified", (name, i)
                if i > 10:
                    break
    finally:
        wb.close()
    raise RuntimeError("Unrecognized tracker schema — no 'API Discovery Tracker' sheet and "
                       "no simplified header row (State / Agency / Outcome or How-provided) found.")


def extract(xlsx_path):
    schema, loc = _detect_schema(xlsx_path)
    if schema == "discovery":
        return _extract_discovery(xlsx_path)
    return _extract_simplified(xlsx_path, loc)


def _clean(v, default=""):
    if v is None or (isinstance(v, float) and np.isnan(v)):
        return default
    s = str(v).strip()
    return default if s in ("nan", "") else s


# ---------- discovery schema (original) ----------

DISCOVERY_COLS = {
    "state": "Jurisidction",
    "stateName": "agency_JURISDICTION::AgencyName",
    "agency": "SS Agency (SS INSERTED COLUMN)",
    "priority": "Priority",
    "coverage": "Coverage - Licensing Coverage",
    "access": "Coverage - Official Access Category",
    "paid": "Coverage - Paid File Status",
    "fee": "Coverage - Known Agency Fee",
    "cadence": "Coverage - Suggested Licensing Cadence",
    "indco": "Individual/Company Data Available?",
    "completeness": "Data Completeness Status",
    "discipline": "Coverage - Included Licensing Discipline",
    "decision": "Coverage - Decision Needed",
    "verifiedAsOf": "Coverage - Verified As Of",
}


def _extract_discovery(xlsx_path):
    df = pd.read_excel(xlsx_path, sheet_name=DISCOVERY_SHEET, header=1, dtype=str)
    df.columns = [str(c).strip() for c in df.columns]

    def col(prefix):
        for c in df.columns:
            if c.strip().lower().startswith(prefix.strip().lower()):
                return c
        return None

    resolved = {k: col(v) for k, v in DISCOVERY_COLS.items()}
    jcol = resolved["state"]
    if jcol is None:
        raise RuntimeError("Could not locate the Jurisdiction column; sheet layout may have changed.")
    df = df[df[jcol].notna() & (df[jcol].astype(str).str.strip() != "")]

    records = []
    for _, r in df.iterrows():
        rec = {k: (_clean(r.get(cn)) if cn is not None else "") for k, cn in resolved.items()}
        if not rec.get("agency"):
            rec["agency"] = _clean(r.get(col("Name")))
        records.append(rec)

    paid = [x for x in records if x["paid"] == "Yes"]
    floor, stated, unknown = 0.0, 0, 0
    for x in paid:
        m = re.search(r"\$\s*([\d,]+(?:\.\d+)?)", x["fee"])
        if m:
            floor += float(m.group(1).replace(",", "")); stated += 1
        else:
            unknown += 1

    verified_vals = [x["verifiedAsOf"] for x in records
                     if re.match(r"^\d{4}-\d{2}-\d{2}$", x["verifiedAsOf"] or "")]
    return {
        "records": records,
        "cost": {"paidAgencies": len(paid), "withStatedFee": stated,
                 "unknownPending": unknown, "oneTimeFloor": round(floor)},
        "meta": {"schema": "discovery", "agencies": len(records),
                 "jurisdictions": int(df[jcol].nunique()),
                 "verifiedAsOf": max(verified_vals) if verified_vals else ""},
    }


# ---------- simplified schema (License Data Tracker) ----------

SIMPLIFIED_COLS = {
    "state": "State",
    "agency": "Agency",
    "url": "Agency URL",
    "public": "Is the data publicly available",       # prefix match (header carries " ? (Y/N)")
    "inclStatus": "Includes Status and/or Expiration",  # absent in newer trackers -> blank
    "inclDiscipline": "Includes Disciplinary Action",   # absent in newer trackers -> blank
    "contacted": "Contacted Agency",
    "outcome": "Outcome",
    "mechanism": "How is the data provided?",
    "cost": "Is there a cost",
    "fee": "Fee Amount",
    "hermai": "Hermai",                                 # matches "Hermai?" and "Hermai Recommended? (Y/N)"
}


def _extract_simplified(xlsx_path, loc):
    sheet_name, header_idx = loc
    df = pd.read_excel(xlsx_path, sheet_name=sheet_name, header=header_idx, dtype=str)
    df.columns = [str(c).strip() for c in df.columns]

    def col(prefix):
        for c in df.columns:
            if c.strip().lower().startswith(prefix.strip().lower()):
                return c
        return None

    resolved = {k: col(v) for k, v in SIMPLIFIED_COLS.items()}
    scol = resolved["state"]
    df = df[df[scol].notna() & (df[scol].astype(str).str.strip() != "")]

    records = [{k: (_clean(r.get(cn)) if cn is not None else "") for k, cn in resolved.items()}
               for _, r in df.iterrows()]

    return {
        "records": records,
        "meta": {"schema": "simplified", "agencies": len(records),
                 "jurisdictions": int(df[scol].nunique()) if len(df) else 0,
                 "verifiedAsOf": ""},
    }


# ---------- render ----------

def render(data, build_str):
    tmpl = "_template_new.txt" if data["meta"].get("schema") == "simplified" else "_template.txt"
    template = open(os.path.join(HERE, tmpl), encoding="utf-8").read()
    payload = json.dumps(
        {k: data[k] for k in ("records", "meta", "cost") if k in data},
        separators=(",", ":"),
    )
    return (template
            .replace("__PAYLOAD__", payload)
            .replace("__VERIFIED__", data["meta"].get("verifiedAsOf", "") or "n/a")
            .replace("__BUILD__", build_str))
